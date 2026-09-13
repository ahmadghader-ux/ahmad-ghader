#!/usr/bin/env python3
"""
Rebuilds the 20-candidate Saudi TikTok scout in the same column layout as the
shared reference sheet "NH Batch 1000 - Round 1 (Saudi trainers, in progress)":
#, Niche, Name (AR), Name (EN), TikTok URL, Handle, Followers, Likes,
Est. age, Face on camera, Evidence / notes.

Honesty note baked into the data itself: the reference sheet's "Est. age" and
"Face on camera" columns were filled by a researcher who actually watched each
creator's TikTok videos and visually judged age/on-camera presence. Direct
TikTok access is blocked in this environment (confirmed 403 policy denial),
so those two columns are marked "n/a - not visually verified (TikTok
inaccessible in this research environment)" for every row here rather than
fabricated, unless a later research pass turns up a publicly documented age
from an independent source (news, LinkedIn, etc.) - those get cited directly.
"""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from build_tiktok_scout import CANDIDATES
from build_tiktok_scout_google import account_url

FONT_NAME = "Arial"
NAVY = "1F4E78"
NOT_VERIFIED_AGE = "n/a — not visually verified (TikTok inaccessible in this research environment)"
NOT_VERIFIED_FACE = "NOT VERIFIED — could not view TikTok videos directly (blocked in this environment)"

# Publicly documented ages found via independent web research (not thumbnails).
# Empty unless/until a citable source is found for a given handle.
DOCUMENTED_AGE = {
    # "@handle": ("age or bracket", "source url"),
}


def split_name(name):
    if " / " in name:
        ar, en = name.split(" / ", 1)
        return ar.strip(), en.strip()
    return name.strip(), name.strip()


def parse_followers_likes(followers_field):
    """Reference sheet keeps Followers and Likes as separate numeric columns.
    Our source data has them combined as free text, e.g. '~178.1K (947.9K likes)'.
    Extract both where the text makes it unambiguous; else 'n/a'."""
    if not followers_field or followers_field.strip().lower() in ("n/a", "n/a on tiktok", "n/a found"):
        return "n/a", "n/a"

    def to_number(s):
        s = s.strip()
        mult = 1
        if s.upper().endswith("M"):
            mult = 1_000_000
            s = s[:-1]
        elif s.upper().endswith("K"):
            mult = 1_000
            s = s[:-1]
        try:
            return round(float(s) * mult)
        except ValueError:
            return None

    m_followers = re.search(r"([\d.]+[KM]?)", followers_field)
    followers = to_number(m_followers.group(1)) if m_followers else None

    m_likes = re.search(r"\(([\d.]+[KM]?)\s*likes\)", followers_field, re.IGNORECASE)
    likes = to_number(m_likes.group(1)) if m_likes else None

    return followers if followers is not None else "n/a", likes if likes is not None else "n/a"


def style_header(ws, row, ncols, fill=NAVY):
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=fill)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, min_width=8, max_width=50):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            length = len(str(cell.value))
            widths[col] = min(max(widths.get(col, min_width), length + 2), max_width)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Candidates"

    headers = ["#", "Niche", "Name (AR)", "Name (EN)", "TikTok URL", "Handle",
               "Followers", "Likes", "Est. age", "Face on camera", "Evidence / notes"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, c in enumerate(CANDIDATES, 1):
        ar, en = split_name(c["name"])
        url, _basis = account_url(c)
        followers, likes = parse_followers_likes(c["followers"])

        if c["handle"] in DOCUMENTED_AGE:
            age, age_src = DOCUMENTED_AGE[c["handle"]]
            est_age = f"{age} (documented: {age_src})"
        else:
            est_age = NOT_VERIFIED_AGE

        evidence = c["why_fit"]
        if c["risk_flag"]:
            evidence += f" | Risk: {c['risk_flag']}"
        if c["status"]:
            evidence += f" | {c['status']}"

        ws.append([i, c["category"], ar, en, url or "n/a", c["handle"],
                   followers, likes, est_age, NOT_VERIFIED_FACE, evidence])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        url_cell = row[4]
        if isinstance(url_cell.value, str) and url_cell.value.startswith("http"):
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"

    autosize(ws, max_width=45)

    out = "NH_TikTok_Scout_AgeFormat.xlsx"
    wb.save(out)
    print(f"Wrote {out}: {len(CANDIDATES)} rows.")


if __name__ == "__main__":
    main()
