#!/usr/bin/env python3
"""
Google-Sheets-ready version of the TikTok-confirmed 20-candidate list (the
25-candidate set minus 5 whose sourcing never actually cited TikTok — see
build_tiktok_confirmed.py for the exclusion list/reasons).

Uses text-rendered bar charts (literal block characters in cells) rather than
embedded openpyxl chart objects: embedded charts were confirmed in this
project not to reliably survive Google's xlsx->Sheets auto-conversion (came
through empty in testing — see build_tiktok_final.py). Text bars are plain
cell content, so they render identically in Excel, Sheets, and LibreOffice.
"""
from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from build_tiktok_scout import TOP_PICKS
from build_tiktok_confirmed import TIKTOK_CANDIDATES, NON_TIKTOK_HANDLES
from build_tiktok_scout_google import account_url
from build_tiktok_final import parse_followers

FONT_NAME = "Arial"
NAVY = "1F4E78"
BAR_COLORS = {"E-commerce": "2E7D32", "Trading": "C62828", "Real Estate": "1565C0"}

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols, fill=NAVY):
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=fill)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, min_width=8, max_width=48):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            length = len(str(cell.value))
            widths[col] = min(max(widths.get(col, min_width), length + 2), max_width)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def main():
    wb = openpyxl.Workbook()

    # ---------------- Creators tab ----------------
    ws = wb.active
    ws.title = "Creators"
    headers = ["#", "Name (AR/EN)", "Niche", "TikTok Account (click)", "Followers",
               "Content Focus", "Why Notable/Fit", "Risk/Flag", "Status"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    for i, c in enumerate(TIKTOK_CANDIDATES, 1):
        url, _ = account_url(c)
        ws.append([i, c["name"], c["category"], c["handle"], c["followers"],
                   c["niche"], c["why_fit"], c["risk_flag"], c["status"]])
        r = ws.max_row
        if url:
            cell = ws.cell(row=r, column=4)
            cell.hyperlink = url
            cell.style = "Hyperlink"

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws, max_width=42)

    # ---------------- Top Picks tab (drop any pick that was removed) ----------------
    ws_top = wb.create_sheet("Top Picks")
    ws_top.append(["Niche", "TikTok Account (click)", "Rank Note"])
    style_header(ws_top, 1, 3)
    ws_top.freeze_panes = "A2"
    by_handle = {c["handle"]: c for c in TIKTOK_CANDIDATES}
    for handle, cat, note in TOP_PICKS:
        if handle in NON_TIKTOK_HANDLES:
            continue
        c = by_handle.get(handle)
        url, _ = account_url(c) if c else (None, None)
        ws_top.append([cat, handle, note])
        r = ws_top.max_row
        if url:
            cell = ws_top.cell(row=r, column=2)
            cell.hyperlink = url
            cell.style = "Hyperlink"
    for row in ws_top.iter_rows(min_row=2, max_row=ws_top.max_row, max_col=3):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws_top, max_width=70)

    # ---------------- Chart tab (text-rendered bars) ----------------
    ws2 = wb.create_sheet("Chart")
    ws2["A1"] = "Candidates by Niche (20 TikTok-confirmed)"
    ws2["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["Niche", "Count", "Bar"])
    style_header(ws2, ws2.max_row, 3)

    counts = Counter(c["category"] for c in TIKTOK_CANDIDATES)
    niches = ["E-commerce", "Trading", "Real Estate"]
    BAR_UNIT = 3
    for niche in niches:
        n = counts.get(niche, 0)
        bar = "█" * (n * BAR_UNIT)
        ws2.append([niche, n, bar])
        r = ws2.max_row
        ws2.cell(row=r, column=3).font = Font(name="Consolas", color=BAR_COLORS.get(niche, "444444"), size=12)

    ws2.append([])
    r = ws2.max_row + 1
    ws2.cell(row=r, column=1, value="Disclosed TikTok Followers (thousands)").font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["Handle", "Niche", "Followers (K)", "Bar"])
    style_header(ws2, ws2.max_row, 4)

    follower_rows = []
    for c in TIKTOK_CANDIDATES:
        f = parse_followers(c["followers"])
        if f is not None:
            follower_rows.append((c["handle"], c["category"], f / 1000))
    follower_rows.sort(key=lambda x: -x[2])

    if follower_rows:
        max_k = max(v for _, _, v in follower_rows)
        for handle, cat, k in follower_rows:
            bar_len = max(1, round((k / max_k) * 40))
            bar = "█" * bar_len
            ws2.append([handle, cat, round(k, 1), bar])
            r = ws2.max_row
            ws2.cell(row=r, column=4).font = Font(name="Consolas", color=BAR_COLORS.get(cat, "444444"), size=12)

    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=4):
        for cell in row:
            if cell.value is not None and cell.column != 3:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 12
    ws2.column_dimensions["D"].width = 45

    # ---------------- Removed tab (transparency on what was filtered out) ----------------
    from build_tiktok_scout import CANDIDATES
    ws3 = wb.create_sheet("Removed (non-TikTok)")
    ws3["A1"] = "Removed from the original 25 — sourced from non-TikTok evidence only"
    ws3["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    ws3.append([])
    ws3.append(["Name (AR/EN)", "Niche", "Handle", "Followers", "Why removed (source basis)"])
    style_header(ws3, ws3.max_row, 5)
    for c in CANDIDATES:
        if c["handle"] in NON_TIKTOK_HANDLES:
            ws3.append([c["name"], c["category"], c["handle"], c["followers"], c["sources"]])
    for row in ws3.iter_rows(min_row=1, max_row=ws3.max_row, max_col=5):
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws3, max_width=55)

    for name in ("Creators", "Top Picks", "Chart", "Removed (non-TikTok)"):
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.font.name != "Consolas":
                    existing = cell.font
                    cell.font = Font(name=FONT_NAME, bold=existing.bold, color=existing.color,
                                      size=existing.size if existing.size else 10)

    out = "NH_TikTok_Scout_ForGoogleSheets.xlsx"
    wb.save(out)
    print(f"Wrote {out}: {len(TIKTOK_CANDIDATES)} TikTok-confirmed candidates, "
          f"{len(follower_rows)} with known followers, "
          f"{len(CANDIDATES) - len(TIKTOK_CANDIDATES)} removed as non-TikTok.")


if __name__ == "__main__":
    main()
