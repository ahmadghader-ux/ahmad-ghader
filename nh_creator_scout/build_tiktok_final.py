#!/usr/bin/env python3
"""
Final consolidated Saudi TikTok scout (25 candidates: original 20 + 5
stricter-bar replacements). Builds two tabs:

  - "Creators": full candidate table with clickable TikTok links.
  - "Chart": a text-rendered bar chart (literal block characters in cells,
    not an embedded Excel chart object). Embedded openpyxl chart objects do
    not reliably survive Google's xlsx->Sheets auto-conversion (confirmed by
    testing in this project — charts came through empty). Text bars are
    just cell content, so they render identically in Excel, Sheets, and
    LibreOffice regardless of the converter.
"""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from build_tiktok_scout import CANDIDATES
from build_tiktok_scout_google import account_url

FONT_NAME = "Arial"
NAVY = "1F4E78"
BAR_COLORS = {"E-commerce": "2E7D32", "Trading": "C62828", "Real Estate": "1565C0"}

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols, fill=NAVY):
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=fill)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, min_width=8, max_width=48):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            length = len(str(cell.value))
            widths[col] = min(max(widths.get(col, min_width), length + 2), max_width)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def parse_followers(followers_field):
    if not followers_field:
        return None
    s = followers_field.strip()
    if s.lower() in ("n/a", "n/a on tiktok", "n/a found", ""):
        return None
    m = re.search(r"([\d.]+)\s*([KM])", s, re.IGNORECASE)
    if not m:
        return None
    num = float(m.group(1))
    mult = 1_000_000 if m.group(2).upper() == "M" else 1_000
    return num * mult


def main():
    wb = openpyxl.Workbook()

    # ---------------- Creators tab ----------------
    ws = wb.active
    ws.title = "Creators"
    headers = ["#", "Name (AR/EN)", "Niche", "TikTok Account (click)", "Followers",
               "Content Focus", "Why Notable/Fit", "Risk/Flag", "Status"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    for i, c in enumerate(CANDIDATES, 1):
        url, _ = account_url(c)
        ws.append([i, c["name"], c["category"], c["handle"], c["followers"],
                   c["niche"], c["why_fit"], c["risk_flag"], c["status"]])
        r = ws.max_row
        if url:
            cell = ws.cell(row=r, column=4)
            cell.hyperlink = url
            cell.style = "Hyperlink"

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws, max_width=42)

    # ---------------- Chart tab (text-rendered bars) ----------------
    ws2 = wb.create_sheet("Chart")
    ws2["A1"] = "Candidates by Niche"
    ws2["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["Niche", "Count", "Bar"])
    style_header(ws2, ws2.max_row, 3)

    from collections import Counter
    counts = Counter(c["category"] for c in CANDIDATES)
    niches = ["E-commerce", "Trading", "Real Estate"]
    max_count = max(counts.values())
    BAR_UNIT = 3  # blocks per unit count
    for niche in niches:
        n = counts.get(niche, 0)
        bar = "█" * (n * BAR_UNIT)
        ws2.append([niche, n, bar])
        r = ws2.max_row
        bar_cell = ws2.cell(row=r, column=3)
        bar_cell.font = Font(name="Consolas", color=BAR_COLORS.get(niche, "444444"), size=12)

    ws2.append([])
    r = ws2.max_row + 1
    ws2.cell(row=r, column=1, value="Disclosed TikTok Followers (thousands)").font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["Handle", "Niche", "Followers (K)", "Bar"])
    style_header(ws2, ws2.max_row, 4)

    follower_rows = []
    for c in CANDIDATES:
        f = parse_followers(c["followers"])
        if f is not None:
            follower_rows.append((c["handle"], c["category"], f / 1000))
    follower_rows.sort(key=lambda x: -x[2])

    if follower_rows:
        max_k = max(v for _, _, v in follower_rows)
        for handle, cat, k in follower_rows:
            bar_len = max(1, round((k / max_k) * 40))
            bar = "█" * bar_len
            ws2.append([handle, cat, round(k, 1), bar])
            r = ws2.max_row
            bar_cell = ws2.cell(row=r, column=4)
            bar_cell.font = Font(name="Consolas", color=BAR_COLORS.get(cat, "444444"), size=12)

    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=4):
        for cell in row:
            if cell.value is not None and cell.column != 3:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 12
    ws2.column_dimensions["D"].width = 45

    for name in ("Creators", "Chart"):
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.font.name != "Consolas":
                    existing = cell.font
                    cell.font = Font(name=FONT_NAME, bold=existing.bold, color=existing.color,
                                      size=existing.size if existing.size else 10)

    out = "NH_TikTok_Scout_Final25.xlsx"
    wb.save(out)
    print(f"Wrote {out}: {len(CANDIDATES)} candidates, {len(follower_rows)} with known followers.")


if __name__ == "__main__":
    main()
