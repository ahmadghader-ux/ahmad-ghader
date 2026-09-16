#!/usr/bin/env python3
"""
Builds NH_TikTok_Scout_Batch2_EcommerceTrading.xlsx — the second Saudi
TikTok creator batch: 24 NEW candidates (13 e-commerce + 11 trading),
excluding every handle already used in the first batch
(build_tiktok_scout.py). Real estate is out of scope for this batch.

Tabs:
  - "New Candidates": full candidate table with clickable TikTok links.
  - "Excluded & Flags": leads considered and rejected during research.
  - "Method": sourcing rules, counts, and caveats (including the
    24-of-25 shortfall note).
"""
import re

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from batch2_candidates import CANDIDATES2, EXCLUDED2, METHOD_NOTES2

FONT_NAME = "Arial"
NAVY = "1F4E78"

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

TIKTOK_URL_RE = re.compile(r"https?://(?:www\.)?tiktok\.com/@[A-Za-z0-9_.]+")
HANDLE_RE = re.compile(r"@[A-Za-z0-9_.]+")


def account_url(candidate):
    m = TIKTOK_URL_RE.search(candidate["sources"])
    if m:
        return m.group(0)
    h = HANDLE_RE.search(candidate["handle"])
    if h:
        return f"https://www.tiktok.com/{h.group(0)}"
    return None


def style_header(ws, row, ncols, fill=NAVY):
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=fill)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, min_width=10, max_width=48):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            length = len(str(cell.value))
            widths[col] = min(max(widths.get(col, min_width), length + 2), max_width)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def set_default_font(ws):
    for row in ws.iter_rows():
        for cell in row:
            existing = cell.font
            cell.font = Font(
                name=FONT_NAME,
                bold=existing.bold if existing else False,
                color=existing.color if existing else None,
                size=existing.size if existing and existing.size else 10,
            )


def main():
    wb = openpyxl.Workbook()

    # ---------------- New Candidates tab ----------------
    ws = wb.active
    ws.title = "New Candidates"
    headers = ["#", "Name (AR/EN)", "Niche", "TikTok Account (click)", "Followers",
               "Content Focus", "Credentials/Background", "Other Socials",
               "Why Notable/Fit", "Risk/Compliance Flag", "Verification Confidence",
               "Status / Flags", "Source URL(s)"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    for i, c in enumerate(CANDIDATES2, 1):
        url = account_url(c)
        ws.append([i, c["name"], c["category"], c["handle"], c["followers"],
                   c["niche"], c["credentials"], c["socials"], c["why_fit"],
                   c["risk_flag"], c["depth_confidence"], c["status"], c["sources"]])
        r = ws.max_row
        if url:
            cell = ws.cell(row=r, column=4)
            cell.hyperlink = url
            cell.style = "Hyperlink"

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws, max_width=42)

    # ---------------- Excluded & Flags tab ----------------
    ws2 = wb.create_sheet("Excluded & Flags")
    ws2["A1"] = "Candidates Excluded During Research (namesake traps, wrong niche, unverifiable)"
    ws2["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["Name", "Niche", "Reason Excluded"])
    style_header(ws2, ws2.max_row, 3)
    for e in EXCLUDED2:
        ws2.append([e["name"], e["category"], e["reason"]])
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=3):
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws2, max_width=60)

    # ---------------- Method tab ----------------
    ws3 = wb.create_sheet("Method")
    ws3["A1"] = "Method, Rules & Counts — Batch 2 (Trading + E-commerce only)"
    ws3["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    ws3.append([])
    ws3.append(["Metric", "Value"])
    style_header(ws3, row=ws3.max_row, ncols=2)
    ws3.append(["Total new candidates", len(CANDIDATES2)])
    ws3.append(["E-commerce", sum(1 for c in CANDIDATES2 if c["category"] == "E-commerce")])
    ws3.append(["Trading", sum(1 for c in CANDIDATES2 if c["category"] == "Trading")])
    ws3.append(["Requested", 25])
    ws3.append(["Shortfall (see notes below)", 25 - len(CANDIDATES2)])
    ws3.append(["Explicitly excluded during research", len(EXCLUDED2)])
    ws3.append([])
    r = ws3.max_row + 1
    ws3.cell(row=r, column=1, value="Notes").font = Font(name=FONT_NAME, bold=True, size=12)
    ws3.append([])
    for note in METHOD_NOTES2:
        ws3.append([note])
        ws3.cell(row=ws3.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws3.merge_cells(start_row=ws3.max_row, start_column=1, end_row=ws3.max_row, end_column=6)
    ws3.column_dimensions["A"].width = 110
    ws3.column_dimensions["B"].width = 14
    for row in ws3.iter_rows(min_row=1, max_row=8, max_col=2):
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    for name in ("New Candidates", "Excluded & Flags", "Method"):
        set_default_font(wb[name])

    out_path = "NH_TikTok_Scout_Batch2_EcommerceTrading.xlsx"
    wb.save(out_path)
    print(f"Wrote {out_path}: {len(CANDIDATES2)} new candidates "
          f"({sum(1 for c in CANDIDATES2 if c['category'] == 'E-commerce')} e-commerce, "
          f"{sum(1 for c in CANDIDATES2 if c['category'] == 'Trading')} trading), "
          f"{len(EXCLUDED2)} excluded.")


if __name__ == "__main__":
    main()
