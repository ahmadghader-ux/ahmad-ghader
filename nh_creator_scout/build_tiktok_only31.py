#!/usr/bin/env python3
"""
Builds NH_TikTok_Only_Creators31.xlsx — every candidate from this project's
two e-commerce/trading research rounds (build_ecom_trading25.py's 25 +
build_ecom_trading25_new.py's 25 = 50) whose confirmed platform is TikTok
specifically, dropping every candidate whose only found account was
Instagram, X, Snapchat, YouTube, or a plain website. 31 of the 50 qualify.

Each row carries a direct, clickable TikTok profile link
(tiktok.com/@handle), derived deterministically from the verified handle
or lifted directly from a cited source URL.
"""
import re
from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from build_ecom_trading25 import CANDIDATES as ROUND1
from build_ecom_trading25_new import CANDIDATES as ROUND2

FONT_NAME = "Arial"
NAVY = "1F4E78"
BAR_COLORS = {"E-commerce": "2E7D32", "Trading": "C62828"}

CANDIDATES = [c for c in (ROUND1 + ROUND2) if c.get("platform") == "tiktok"]

TIKTOK_URL_RE = re.compile(r"https?://(?:www\.)?tiktok\.com/@[A-Za-z0-9_.]+")
HANDLE_RE = re.compile(r"@[A-Za-z0-9_.]+")


def tiktok_url(c):
    m = TIKTOK_URL_RE.search(c["sources"])
    if m:
        return m.group(0)
    h = HANDLE_RE.search(c["handle"])
    return f"https://www.tiktok.com/{h.group(0)}" if h else None


def style_header(ws, row, ncols, fill=NAVY):
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=fill)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, min_width=8, max_width=48):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            length = len(str(cell.value))
            widths[col] = min(max(widths.get(col, min_width), length + 2), max_width)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def parse_followers(followers_field):
    if not followers_field:
        return None
    s = followers_field.strip()
    if s.lower().startswith("n/a") or s == "":
        return None
    m = re.search(r"([\d.]+)\s*([KM])", s, re.IGNORECASE)
    if not m:
        return None
    num = float(m.group(1))
    mult = 1_000_000 if m.group(2).upper() == "M" else 1_000
    return num * mult


def main():
    counts = Counter(c["category"] for c in CANDIDATES)
    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ---------------- Creators tab ----------------
    ws = wb.active
    ws.title = "Creators"
    headers = ["#", "Name (AR/EN)", "Niche", "TikTok Profile (click)", "Followers",
               "Content Focus", "Why Notable/Fit", "Risk/Flag", "Status"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    for i, c in enumerate(CANDIDATES, 1):
        url = tiktok_url(c)
        ws.append([i, c["name"], c["category"], url or "n/a", c["followers"],
                   c["niche"], c["why_fit"], c["risk_flag"], c["status"]])
        r = ws.max_row
        if url:
            cell = ws.cell(row=r, column=4)
            cell.hyperlink = url
            cell.style = "Hyperlink"

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws, max_width=42)

    # ---------------- Chart tab ----------------
    ws2 = wb.create_sheet("Chart")
    ws2["A1"] = f"{len(CANDIDATES)} TikTok-Only Creators by Niche"
    ws2["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["Niche", "Count", "Bar"])
    style_header(ws2, ws2.max_row, 3)

    BAR_UNIT = 3
    for niche in ["E-commerce", "Trading"]:
        n = counts.get(niche, 0)
        bar = "█" * (n * BAR_UNIT)
        ws2.append([niche, n, bar])
        r = ws2.max_row
        ws2.cell(row=r, column=3).font = Font(name="Consolas", color=BAR_COLORS.get(niche, "444444"), size=12)

    ws2.append([])
    ws2.cell(row=ws2.max_row + 1, column=1, value="Disclosed TikTok Followers (thousands)").font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["Handle", "Niche", "Followers (K)", "Bar"])
    style_header(ws2, ws2.max_row, 4)

    follower_rows = []
    for c in CANDIDATES:
        f = parse_followers(c["followers"])
        if f is not None:
            follower_rows.append((c["handle"], c["category"], f / 1000))
    follower_rows.sort(key=lambda x: -x[2])

    if follower_rows:
        max_k = max(v for _, _, v in follower_rows)
        for handle, cat, k in follower_rows:
            bar_len = max(1, round((k / max_k) * 40))
            bar = "█" * bar_len
            ws2.append([handle, cat, round(k, 1), bar])
            r = ws2.max_row
            ws2.cell(row=r, column=4).font = Font(name="Consolas", color=BAR_COLORS.get(cat, "444444"), size=12)

    ws2.append([])
    r = ws2.max_row + 1
    ws2.cell(row=r, column=1, value="All TikTok Profiles").font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["#", "Name", "Niche", "TikTok Profile Link"])
    style_header(ws2, ws2.max_row, 4)
    for i, c in enumerate(CANDIDATES, 1):
        url = tiktok_url(c)
        ws2.append([i, c["name"], c["category"], url or "n/a"])
        r = ws2.max_row
        if url:
            cell = ws2.cell(row=r, column=4)
            cell.hyperlink = url
            cell.style = "Hyperlink"

    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=4):
        for cell in row:
            if cell.value is not None and cell.column != 3:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 12
    ws2.column_dimensions["D"].width = 45

    for name in ("Creators", "Chart"):
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.font.name != "Consolas":
                    existing = cell.font
                    cell.font = Font(name=FONT_NAME, bold=existing.bold, color=existing.color,
                                      size=existing.size if existing.size else 10)

    out = "NH_TikTok_Only_Creators31.xlsx"
    wb.save(out)
    print(f"Wrote {out}: {len(CANDIDATES)} TikTok creators "
          f"({counts['E-commerce']} E-commerce, {counts['Trading']} Trading), "
          f"{len(follower_rows)} with a disclosed follower count.")


if __name__ == "__main__":
    main()
