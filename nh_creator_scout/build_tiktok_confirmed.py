#!/usr/bin/env python3
"""
TikTok-only filter of the 25-candidate list: drops candidates whose research
sourcing never actually cited TikTok (a TikTok URL, TikTok video/profile
snippet, or TikTok search result) — only Instagram, Snapchat, a personal
website, or a third-party aggregator listing. Those 5 were "found" via
non-TikTok evidence and their TikTok presence was only ever a guessed
tiktok.com/@handle, never confirmed.

Excluded (non-TikTok-sourced):
  @malazmarketing    — sourced only from malazmarketing.com (no TikTok mention)
  @katibschool       — sourced from Instagram + own site; followers explicitly "n/a on TikTok"
  @fuad8k            — sourced from fuadacademy.com + X; no TikTok mention
  @kii31t            — sourced from search snippets + Snapchat; followers "n/a on TikTok"
  @deserttrader6906  — sourced from a third-party Heepsy aggregator listing, not TikTok itself

Remaining: 20 candidates, all with at least one TikTok-specific citation
(a tiktok.com URL, or an explicit "TikTok video/profile/search" source note).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from build_tiktok_scout import CANDIDATES
from build_tiktok_scout_google import account_url

FONT_NAME = "Arial"
NAVY = "1F4E78"

NON_TIKTOK_HANDLES = {
    "@malazmarketing",
    "@katibschool",
    "@fuad8k",
    "@kii31t",
    "@deserttrader6906",
}

TIKTOK_CANDIDATES = [c for c in CANDIDATES if c["handle"] not in NON_TIKTOK_HANDLES]

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols, fill=NAVY):
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=fill)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, min_width=8, max_width=48):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            length = len(str(cell.value))
            widths[col] = min(max(widths.get(col, min_width), length + 2), max_width)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Creators"
    headers = ["#", "Name (AR/EN)", "Niche", "TikTok Account (click)", "Followers",
               "Content Focus", "Why Notable/Fit", "Risk/Flag", "Status"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    for i, c in enumerate(TIKTOK_CANDIDATES, 1):
        url, _ = account_url(c)
        ws.append([i, c["name"], c["category"], c["handle"], c["followers"],
                   c["niche"], c["why_fit"], c["risk_flag"], c["status"]])
        r = ws.max_row
        if url:
            cell = ws.cell(row=r, column=4)
            cell.hyperlink = url
            cell.style = "Hyperlink"

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws, max_width=42)

    ws2 = wb.create_sheet("Removed (non-TikTok)")
    ws2["A1"] = "Removed — sourced from non-TikTok evidence only"
    ws2["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    ws2.append([])
    ws2.append(["Name (AR/EN)", "Niche", "Handle", "Followers", "Why removed (source basis)"])
    style_header(ws2, ws2.max_row, 5)
    for c in CANDIDATES:
        if c["handle"] in NON_TIKTOK_HANDLES:
            ws2.append([c["name"], c["category"], c["handle"], c["followers"], c["sources"]])
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=5):
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws2, max_width=55)

    for name in ("Creators", "Removed (non-TikTok)"):
        for row in wb[name].iter_rows():
            for cell in row:
                existing = cell.font
                cell.font = Font(name=FONT_NAME, bold=existing.bold, color=existing.color,
                                  size=existing.size if existing.size else 10)

    out = "NH_TikTok_Scout_Confirmed20.xlsx"
    wb.save(out)
    print(f"Wrote {out}: {len(TIKTOK_CANDIDATES)} confirmed TikTok candidates "
          f"(removed {len(CANDIDATES) - len(TIKTOK_CANDIDATES)} non-TikTok-sourced).")


if __name__ == "__main__":
    main()
